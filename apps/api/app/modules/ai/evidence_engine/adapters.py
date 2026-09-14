"""Local, immutable source adapters for the evidence-first experiment."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .models import SourceDocument, SourceFact, SourceSection

_GENERIC_MATRIX_HEADERS = {
    "поле",
    "характеристика",
    "параметр",
    "свойство",
    "field",
    "attribute",
    "parameter",
    "property",
}
_CATALOG_HEADERS = {
    "артикул",
    "sku",
    "наименование",
    "название",
    "weight",
    "вес",
    "размеры",
    "описание",
}


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).replace("\u00a0", " ").split())


def _slug(value: str) -> str:
    normalized = re.sub(r"[^\w]+", "-", value.casefold(), flags=re.UNICODE).strip("-")
    return normalized[:80] or "item"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _worksheet_role(rows: list[list[str]]) -> str:
    if not rows:
        return "supporting"
    headers = {cell.casefold() for cell in rows[0] if cell}
    first_header = rows[0][0].casefold() if rows[0] else ""
    if first_header in _GENERIC_MATRIX_HEADERS and len([x for x in rows[0][1:] if x]) >= 2:
        return "primary"
    catalog_hits = len(headers & _CATALOG_HEADERS)
    if catalog_hits >= 2 or len(rows) > 100:
        return "supporting"
    return "primary"


class XlsxEvidenceAdapter:
    """Preserve matrix ownership and keep catalogs out of curriculum planning."""

    def read(self, path: Path) -> SourceDocument:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: list[SourceSection] = []
        try:
            for worksheet in workbook.worksheets:
                rows = [
                    [_text(cell) for cell in row]
                    for row in worksheet.iter_rows(values_only=True)
                ]
                rows = [row for row in rows if any(row)]
                role = _worksheet_role(rows)
                facts = (
                    self._matrix_facts(worksheet.title, rows)
                    if self._is_matrix(rows)
                    else self._row_facts(worksheet.title, rows)
                )
                sections.append(
                    SourceSection(
                        section_id=f"sheet:{_slug(worksheet.title)}",
                        title=worksheet.title,
                        role=role,  # type: ignore[arg-type]
                        facts=tuple(facts),
                    )
                )
        finally:
            workbook.close()
        return SourceDocument(
            source_id=f"sha256:{_sha256(path)}",
            title=path.stem,
            kind="spreadsheet",
            sections=tuple(sections),
            source_sha256=_sha256(path),
        )

    @staticmethod
    def _is_matrix(rows: list[list[str]]) -> bool:
        return bool(
            rows
            and rows[0]
            and rows[0][0].casefold() in _GENERIC_MATRIX_HEADERS
            and len([x for x in rows[0][1:] if x]) >= 2
        )

    @staticmethod
    def _matrix_facts(sheet: str, rows: list[list[str]]) -> list[SourceFact]:
        if not rows:
            return []
        entities = rows[0][1:]
        facts: list[SourceFact] = []
        for row_number, row in enumerate(rows[1:], start=2):
            attribute = row[0] if row else ""
            if not attribute:
                continue
            for column_number, entity in enumerate(entities, start=2):
                value = row[column_number - 1] if column_number - 1 < len(row) else ""
                if not entity or not value:
                    continue
                facts.append(
                    SourceFact(
                        fact_id=f"xlsx:{_slug(sheet)}:r{row_number}:c{column_number}",
                        subject=entity,
                        attribute=attribute,
                        value=value,
                        source_locator=f"sheet={sheet};row={row_number};column={column_number}",
                    )
                )
        return facts

    @staticmethod
    def _row_facts(sheet: str, rows: list[list[str]]) -> list[SourceFact]:
        if len(rows) < 2:
            return []
        headers = rows[0]
        subject_index = next(
            (
                index
                for index, header in enumerate(headers)
                if header.casefold() in {"наименование", "название", "name", "товар"}
            ),
            0,
        )
        facts: list[SourceFact] = []
        for row_number, row in enumerate(rows[1:], start=2):
            subject = row[subject_index] if subject_index < len(row) else ""
            if not subject:
                subject = next((cell for cell in row if cell), f"Строка {row_number}")
            for column_number, header in enumerate(headers, start=1):
                value = row[column_number - 1] if column_number - 1 < len(row) else ""
                if not header or not value or column_number - 1 == subject_index:
                    continue
                facts.append(
                    SourceFact(
                        fact_id=f"xlsx:{_slug(sheet)}:r{row_number}:c{column_number}",
                        subject=subject,
                        attribute=header,
                        value=value,
                        source_locator=f"sheet={sheet};row={row_number};column={column_number}",
                    )
                )
        return facts


def narrative_document_from_pages(
    *,
    path: Path,
    pages: list[str],
) -> SourceDocument:
    """Build clause-level evidence while retaining headings and page provenance."""

    section_data: dict[str, dict[str, object]] = {}
    current_title = "Вводная часть"
    current_role = "supporting"
    paragraph: list[str] = []
    paragraph_page = 1

    def ensure_section(title: str, role: str) -> dict[str, object]:
        return section_data.setdefault(title, {"role": role, "facts": []})

    def flush_paragraph() -> None:
        nonlocal paragraph
        value = " ".join(part.strip() for part in paragraph if part.strip())
        paragraph = []
        if len(value) < 25:
            return
        section = ensure_section(current_title, current_role)
        section_facts = section["facts"]
        assert isinstance(section_facts, list)
        table_like = (
            ("Условие" in value and "Описание" in value)
            or (
                value.count("Максимальн") >= 2
                and len(re.findall(r"\d+(?:[.,]\d+)?", value)) >= 6
            )
        )
        section_facts.append(
            (
                paragraph_page,
                value,
                0.6 if table_like else 0.9,
                "local_ocr_table" if table_like else "local_ocr",
            )
        )

    def is_upper_heading(line: str) -> bool:
        letters = [char for char in line if char.isalpha()]
        return bool(letters) and sum(char.isupper() for char in letters) / len(letters) >= 0.82

    major_heading_re = re.compile(r"^(?:\d+|[Зз])\.\s+.+")
    clause_re = re.compile(r"^(?:\d+|[Зз])\s*[).,]\s+.+")

    for page_number, page_text in enumerate(pages, start=1):
        lines = [" ".join(line.split()) for line in page_text.splitlines() if line.strip()]
        for line in lines:
            if re.fullmatch(r"\d{1,3}", line):
                continue
            if line.casefold().startswith("содержание"):
                flush_paragraph()
                current_title = "Содержание"
                current_role = "supporting"
                ensure_section(current_title, current_role)
                continue
            if line.casefold().startswith("приложение"):
                flush_paragraph()
                current_title = line
                current_role = "supporting"
                ensure_section(current_title, current_role)
                continue
            if major_heading_re.match(line) and is_upper_heading(line):
                flush_paragraph()
                current_title = line
                current_role = "primary"
                ensure_section(current_title, current_role)
                continue
            if current_role == "primary" and not paragraph and is_upper_heading(line):
                # OCR often splits a numbered chapter heading over two visual lines.
                old_title = current_title
                current_title = f"{old_title} {line}"
                previous = section_data.pop(old_title, None)
                if previous is not None:
                    section_data[current_title] = previous
                else:
                    ensure_section(current_title, current_role)
                continue
            if clause_re.match(line):
                flush_paragraph()
                paragraph_page = page_number
                paragraph = [line]
                continue
            if paragraph:
                paragraph.append(line)
            elif current_role == "supporting":
                paragraph_page = page_number
                paragraph = [line]
        # Keep an unfinished numbered clause open across page boundaries.
    flush_paragraph()

    sections: list[SourceSection] = []
    for section_index, (title, data) in enumerate(section_data.items(), start=1):
        raw_facts = data["facts"]
        assert isinstance(raw_facts, list)
        facts = tuple(
            SourceFact(
                fact_id=f"pdf:s{section_index}:f{fact_index}",
                subject=title,
                attribute=_narrative_attribute(value),
                value=value,
                source_locator=f"page={page_number};section={section_index};clause={fact_index}",
                confidence=confidence,
                uncertainty=uncertainty,
            )
            for fact_index, (page_number, value, confidence, uncertainty) in enumerate(
                raw_facts,
                start=1,
            )
        )
        if facts:
            sections.append(
                SourceSection(
                    section_id=f"pdf:s{section_index}",
                    title=title,
                    role=data["role"],  # type: ignore[arg-type]
                    facts=facts,
                )
            )
    digest = _sha256(path)
    return SourceDocument(
        source_id=f"sha256:{digest}",
        title=path.stem,
        kind="narrative",
        sections=tuple(sections),
        source_sha256=digest,
    )


def _narrative_attribute(sentence: str) -> str:
    normalized = sentence.casefold().replace("ё", "е")
    if re.search(r"\b(?:дн|дня|дней|месяц|месяца|месяцев|час|часов|срок)\w*\b", normalized):
        return "срок"
    if re.search(r"\b(?:процент|ставк|вознагражден|тенге|сумм)\w*\b", normalized):
        return "финансовое условие"
    if re.search(r"\b(?:вправе|право)\w*\b", normalized):
        return "право"
    if re.search(r"\b(?:обязан|должен|необходимо)\w*\b", normalized):
        return "обязанность"
    if re.search(r"\b(?:запрещен|не допускается|не вправе)\w*\b", normalized):
        return "запрет"
    return "положение"
