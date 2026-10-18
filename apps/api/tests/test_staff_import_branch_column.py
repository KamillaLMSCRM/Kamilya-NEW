from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from openpyxl import Workbook

from app.modules.staff_import_legacy_adapter import LegacyColumnMap, adapt_legacy_rows
from app.modules.staff_import_sessions import ImportMode
from app.modules.users.staff_import_service import parse_upload


def _xlsx_with_branch_column(*, merged: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сотрудники"
    sheet.append(["Табельный номер", "ФИО", "Должность", "Филиал"])
    sheet.append(["001", "Иванов Иван", "Кассир", "Павлодар"])
    sheet.append(["002", "Петров Пётр", "Менеджер", None])
    if merged:
        sheet.merge_cells("D2:D3")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_branch_alias_and_forward_fill_are_preserved_in_parsed_rows():
    parsed = parse_upload("staff.xlsx", _xlsx_with_branch_column())

    assert parsed.missing_required_columns == []
    assert parsed.detected_columns["Филиал"] == "branch"
    assert [row.branch for row in parsed.rows] == ["Павлодар", "Павлодар"]
    assert [row.department for row in parsed.rows] == ["", ""]
    assert parsed.invalid_rows == []


def test_merged_branch_cells_are_expanded_before_parsing():
    parsed = parse_upload("staff.xlsx", _xlsx_with_branch_column(merged=True))

    assert parsed.invalid_rows == []
    assert [row.branch for row in parsed.rows] == ["Павлодар", "Павлодар"]


def test_six_column_layout_proposal_refs_follow_physical_branch_and_department_cells():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сотрудники"
    sheet.append(["Табельный номер", "ФИО", "Должность", "Филиал", "Отдел", "Email"])
    sheet.append(["001", "Тестов Первый", "Оператор", "Филиал А", "Отдел А", None])
    output = BytesIO()
    workbook.save(output)

    parsed = parse_upload("staff.xlsx", output.getvalue())
    proposal = adapt_legacy_rows(
        tenant_id=uuid4(),
        source_file_name="staff.xlsx",
        source_file_sha256="a" * 64,
        rows=parsed.rows,
        sheet_name=parsed.sheet_name or "Лист1",
        columns=LegacyColumnMap.from_headers(parsed.raw_columns, parsed.detected_columns),
        mode=ImportMode.ADD_OR_UPDATE,
    )

    assert proposal.branches[0].source_refs[0].column == "D"
    assert proposal.departments[0].source_refs[0].column == "E"
    assert proposal.positions[0].source_refs[0].column == "C"
    assert [ref.column for ref in proposal.staff[0].source_refs] == ["A", "B", "D"]


def test_reordered_csv_manual_mapping_tracks_split_names_and_ignores_unmapped_column():
    content = (
        b"Ignored,Employee Code,Family,Given,Unused,Job,Location,Unit\n"
        b"x,009,Testov,First,x,Consultant,Branch A,Dept A\n"
    )
    mapping = {
        "personnel_number": "Employee Code",
        "last_name": "Family",
        "first_name": "Given",
        "position": "Job",
        "branch": "Location",
        "department": "Unit",
    }

    parsed = parse_upload("staff.csv", content, mapping=mapping)
    proposal = adapt_legacy_rows(
        tenant_id=uuid4(),
        source_file_name="staff.csv",
        source_file_sha256="b" * 64,
        rows=parsed.rows,
        columns=LegacyColumnMap.from_headers(parsed.raw_columns, parsed.detected_columns),
    )

    assert parsed.missing_required_columns == []
    assert (proposal.staff[0].first_name, proposal.staff[0].last_name) == ("First", "Testov")
    assert proposal.branches[0].source_refs[0].column == "G"
    assert proposal.departments[0].source_refs[0].column == "H"
    assert proposal.positions[0].source_refs[0].column == "F"
    assert [ref.column for ref in proposal.staff[0].source_refs] == ["B", "D", "C", "G"]
    assert "E" not in {ref.column for ref in proposal.staff[0].source_refs}


def test_reordered_xlsx_auto_mapping_tracks_coordinates_and_omits_absent_optional_refs():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сотрудники"
    sheet.append(["Отдел", "Комментарий", "Филиал", "Email", "Должность", "Табельный номер", "ФИО"])
    sheet.append(["Отдел А", "ignored", "Филиал А", None, "Оператор", "010", "Тестов Первый"])
    output = BytesIO()
    workbook.save(output)

    parsed = parse_upload("staff.xlsx", output.getvalue())
    proposal = adapt_legacy_rows(
        tenant_id=uuid4(),
        source_file_name="staff.xlsx",
        source_file_sha256="c" * 64,
        rows=parsed.rows,
        columns=LegacyColumnMap.from_headers(parsed.raw_columns, parsed.detected_columns),
    )

    assert parsed.missing_required_columns == []
    assert proposal.branches[0].source_refs[0].column == "C"
    assert proposal.departments[0].source_refs[0].column == "A"
    assert proposal.positions[0].source_refs[0].column == "E"
    assert [ref.column for ref in proposal.staff[0].source_refs] == ["F", "G", "C"]
    assert all(
        ref.column != "D"
        for item in (*proposal.branches, *proposal.departments, *proposal.positions, *proposal.staff)
        for ref in item.source_refs
    )
