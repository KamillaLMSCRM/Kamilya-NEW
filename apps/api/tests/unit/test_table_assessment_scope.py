"""Exercise table scope through real DOCX conversion, chunking and admission."""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from app.modules.ai.direct_source import DirectSourceChunk, DirectSourceCorpus, DirectSourceDocument
from app.modules.ai.evidence_engine.application import build_evidence_source
from app.modules.ai.evidence_engine.assessment_axes import (
    derive_assessment_axes,
    has_partial_unatomized_source_contract,
)
from app.modules.ai.evidence_engine.models import LessonDraft
from app.modules.ai.evidence_engine.semantic_assessment import generate_block_assessment
from app.modules.ai.ingestion import DocumentChunker, _local_convert


async def _converted_axes(tmp_path: Path, value: str):
    path = tmp_path / "synthetic-approval.docx"
    document = Document()
    document.add_heading("Порядок согласования операций", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Категория"
    table.cell(0, 1).text = "Порядок согласования"
    table.cell(1, 0).text = "Категория Б"
    table.cell(1, 1).text = value
    document.save(path)
    converted = await _local_convert(str(path))
    rows = DocumentChunker().chunk_markdown(converted["markdown"], "synthetic", path.name)
    chunks = tuple(DirectSourceChunk(
        chunk_id=f"synthetic:{index}", doc_id="synthetic", doc_name=path.name,
        title="Порядок согласования операций", text=row["text"],
        headings=tuple(json.loads(row["metadata"]["headings"])),
        source_revision="synthetic-v1", chunk_index=index,
        table_fragment=row["metadata"].get("worksheet_table_fragment"),
    ) for index, row in enumerate(rows))
    corpus = DirectSourceCorpus("synthetic", (DirectSourceDocument(
        "synthetic", "Порядок согласования операций", path.name,
        "training_material", "synthetic-v1", chunks,
    ),), len(converted["markdown"]), len(chunks))
    source = build_evidence_source(corpus).document
    facts = [fact for section in source.sections for fact in section.facts]
    lesson = LessonDraft("lesson", "Правила", "Согласование", "Согласовать операцию",
                         "", tuple(f.fact_id for f in facts), (), 3)
    return facts, derive_assessment_axes(lesson, facts, block_id="scope")


@pytest.mark.asyncio
@pytest.mark.parametrize("marker", [")", "."])
async def test_adjacent_exception_deadline_keeps_active_branch_and_trigger(
    tmp_path: Path, marker: str,
) -> None:
    value = (
        "Согласование операций производится ответственным специалистом: "
        f"1{marker} при расположении пункта в пределах города - специалист проверяет документы.\n"
        "При обычной выдаче специалист подписывает документ в день выдачи операции.\n"
        "Если выдача операции производится после 18:00, в выходные или праздничные дни, "
        "то решение должно быть предварительно согласовано по телефону.\n"
        "Согласование операции ответственный специалист производит в первый рабочий "
        "день с момента выдачи операции.\n"
        f"2{marker} при расположении пункта вне города - согласование производится во время выездной проверки."
    )
    facts, axes = await _converted_axes(tmp_path, value)
    deadline = next(axis for axis in axes if axis.correct_value ==
                    "в первый рабочий день с момента выдачи операции")
    assert "при расположении пункта в пределах города" in deadline.required_prompt
    assert "после 18:00, в выходные или праздничные дни" in deadline.required_prompt
    assert "вне города" not in deadline.required_prompt
    assert len(deadline.required_prompt) < 900
    assert deadline.correct_value in next(f.value for f in facts if f.fact_id == deadline.primary_fact_id)


@pytest.mark.asyncio
async def test_refusal_marking_step_does_not_claim_complete_procedure(tmp_path: Path) -> None:
    value = (
        "Ответственный специалист проверяет документы перед согласованием операции.\n"
        "Если специалист не согласовывает выданный микрокредит, то в бумажном журнале и в системе "
        "учета проставляется отметка, что микрокредит НЕ СОГЛАСОВАН.\n"
        "В этом случае составляется акт несоответствия с расчетом суммы.\n"
        "С сотрудника берется объяснительная. Акты предоставляются в отдел контроля."
    )
    facts, axes = await _converted_axes(tmp_path, value)
    marking = next(axis for axis in axes if "НЕ СОГЛАСОВАН" in axis.correct_value)
    assert "Какую отметку" in marking.required_prompt
    assert "в бумажном журнале и в системе учета" in marking.required_prompt
    assert "Если специалист не согласовывает выданный микрокредит" in marking.required_prompt
    assert "Как следует поступить" not in marking.required_prompt
    assert "акт несоответствия" not in marking.required_prompt
    assert marking.correct_value in next(f.value for f in facts if f.fact_id == marking.primary_fact_id)
    assert has_partial_unatomized_source_contract(
        next(f for f in facts if f.fact_id == marking.primary_fact_id)
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix, forbidden, expected", [
    (
        "2) при расположении пункта вне города - согласование операции специалист "
        "производит в течение 24 часов с момента выдачи операции.",
        "после 18:00", "в течение 24 часов с момента выдачи операции",
    ),
    (
        "Согласование операции специалист производит во всех случаях в течение "
        "24 часов с момента выдачи операции.",
        "после 18:00", "в течение 24 часов с момента выдачи операции",
    ),
])
async def test_exception_cannot_cross_branch_or_override_explicit_universal_scope(
    tmp_path: Path, suffix: str, forbidden: str, expected: str,
) -> None:
    value = (
        "Для согласования операций специалист проверяет документы и журнал. "
        "1) при расположении пункта в пределах города - специалист проверяет документы. "
        "Если выдача операции производится после 18:00, то необходимо согласование по телефону. "
        + suffix
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    deadline = next(axis for axis in axes if axis.correct_value == expected)
    assert forbidden not in deadline.required_prompt
    if suffix.startswith("2)"):
        assert "вне города" in deadline.required_prompt
        assert "в пределах города" not in deadline.required_prompt


@pytest.mark.asyncio
@pytest.mark.parametrize("context", [
    "Если доставка посылки производится после 18:00, то необходимо уведомить диспетчера.",
    "Если выдача операции производится после 18:00 и " + "подтверждена сверкой документов, " * 24
    + "то необходимо предварительное согласование по телефону.",
    "2) по отдельному регламенту проводится проверка без описанного условия.",
])
async def test_unresolved_or_overlong_scope_is_not_emitted_as_unconditional_deadline(
    tmp_path: Path, context: str,
) -> None:
    value = (
        "Для согласования операций специалист проверяет документы и журнал. "
        "1) при расположении пункта в пределах города - специалист проверяет документы. "
        + context + " Согласование операции специалист производит в первый рабочий "
        "день с момента выдачи операции."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    assert not any(axis.correct_value == "в первый рабочий день с момента выдачи операции"
                   for axis in axes)


@pytest.mark.asyncio
async def test_notification_step_asks_recipient_and_timing_not_unrelated_act(tmp_path: Path) -> None:
    value = (
        "Перед приемом смены сотрудник проверяет соответствие имущества и документов. "
        "Если залоговое имущество вызывает подозрения в отношении качества, то сменный "
        "сотрудник обязан сразу поставить в известность руководителя отдела. "
        "После получения такой информации руководитель проводит проверку и составляет акт. "
        "Результаты проверки передаются в отдел контроля для отдельного решения."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    notification = next(axis for axis in axes if "обязан сразу поставить" in axis.correct_value)
    assert "Кого и когда" in notification.required_prompt
    assert "сменный сотрудник" in notification.required_prompt
    assert "Если залоговое имущество вызывает подозрения" in notification.required_prompt
    assert "Как следует поступить" not in notification.required_prompt


@pytest.mark.asyncio
async def test_same_actor_and_object_do_not_prove_deadline_condition_binding(tmp_path: Path) -> None:
    value = (
        "Согласование операции включает проверку документов и решение специалиста. "
        "Если специалист не согласовывает операцию, то необходимо уведомить руководителя. "
        "Специалист согласовывает операцию в первый рабочий день. "
        "По окончании операции специалист передает документы в архив для хранения."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    assert not any(a.correct_value == "в первый рабочий день" for a in axes)


@pytest.mark.asyncio
@pytest.mark.parametrize("action", [
    "сотрудники обязаны немедленно уведомить руководителя отдела",
    "сотрудница должна незамедлительно сообщить руководителю отдела",
    "сотрудник проставляет отметку НЕ СОГЛАСОВАН в журнале и в системе учета",
])
async def test_unsupported_step_grammar_never_falls_back_to_whole_workflow(
    tmp_path: Path, action: str,
) -> None:
    value = (
        "Перед приемом смены сотрудник проверяет имущество и документы по перечню. "
        f"Если специалист не согласовывает выданный микрокредит, то {action}. "
        "В этом случае составляется акт несоответствия с расчетом суммы. "
        "С сотрудника берется объяснительная, материалы передаются в отдел контроля."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    for axis in axes:
        if action in axis.correct_value:
            assert "Как следует поступить" not in axis.required_prompt
            assert "какое действие предусмотрено" not in axis.required_prompt


@pytest.mark.asyncio
async def test_equal_deadlines_under_distinct_branches_remain_distinct_targets(tmp_path: Path) -> None:
    value = (
        "Согласование операций производится ответственным специалистом по следующим правилам. "
        "1) при расположении пункта в пределах города - проверка документов обязательна. "
        "Согласование операции производится в течение 24 часов с момента выдачи операции. "
        "2) при расположении пункта вне города - проверка документов обязательна. "
        "Согласование операции производится в течение 24 часов с момента выдачи операции."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    deadlines = [a for a in axes if a.correct_value == "в течение 24 часов с момента выдачи операции"]
    assert len(deadlines) == 2
    assert len({a.axis_id for a in deadlines}) == 2
    assert "в пределах города" in deadlines[0].required_prompt
    assert "вне города" in deadlines[1].required_prompt


@pytest.mark.asyncio
async def test_admitted_refusal_cannot_claim_complete_coverage_of_unassessed_siblings(tmp_path: Path) -> None:
    value = (
        "Сотрудник обязан проверить документ перед согласованием операции. "
        "Если специалист не согласовывает выданный микрокредит, то в журнале и в системе "
        "учета проставляется отметка, что микрокредит НЕ СОГЛАСОВАН. "
        "В этом случае составляется акт несоответствия с расчетом суммы. "
        "С сотрудника берется объяснительная и обязательство. Акты предоставляются в отдел контроля."
    )
    facts, _axes = await _converted_axes(tmp_path, value)
    lesson = LessonDraft("lesson", "Правила", "Согласование", "Согласовать операцию",
                         "", tuple(f.fact_id for f in facts), (), 3)

    class EmptyAuthor:
        async def ainvoke_validated(self, messages, parser, **_kwargs):
            assert json.loads(messages[-1]["content"])["task"] == "assessment_generate"
            return SimpleNamespace(value=parser('{"questions":[]}'), attempt_count=1)

    result = await generate_block_assessment([lesson], {f.fact_id: f for f in facts}, EmptyAuthor())
    assert result.audit["terminal_status"] == "review_required"
    assert result.audit["coverage"]["requires_review"] is True
    assert any(row["reason"] == "partial_unatomized_source_contract"
               for row in result.audit["axis_outcomes"])


@pytest.mark.asyncio
async def test_quarantined_step_cannot_reenter_via_salient_sentence_fallback(tmp_path: Path) -> None:
    value = (
        "Проверку документов специалист производит в течение 24 часов с момента регистрации заявки. "
        "Если специалист не согласовывает выданный микрокредит, сотрудник обязан проставить "
        "отметку НЕ СОГЛАСОВАН в журнале и в системе учета. "
        "В этом случае составляется акт несоответствия с расчетом суммы. "
        "С сотрудника берется объяснительная и обязательство, акты передаются в отдел контроля."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    assert any(a.correct_value == "в течение 24 часов с момента регистрации заявки" for a in axes)
    assert not any("НЕ СОГЛАСОВАН" in a.correct_value for a in axes)


@pytest.mark.asyncio
@pytest.mark.parametrize("marker", [")", "."])
async def test_condition_embedded_after_branch_header_cannot_disappear(
    tmp_path: Path, marker: str,
) -> None:
    value = (
        "Согласование операций производится ответственным специалистом по следующим правилам: "
        f"1{marker} при расположении пункта в пределах города - Если выдача операции производится "
        "после 18:00, в выходные или праздничные дни, то решение необходимо согласовать по телефону. "
        "Согласование операции специалист производит в первый рабочий день с момента выдачи операции. "
        f"2{marker} при расположении пункта вне города - согласование производится при выездной проверке."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    for axis in axes:
        if axis.correct_value == "в первый рабочий день с момента выдачи операции":
            assert "после 18:00, в выходные или праздничные дни" in axis.required_prompt
            assert "в пределах города" in axis.required_prompt


@pytest.mark.asyncio
async def test_conditional_infinitive_cannot_be_mistaken_for_the_required_action(tmp_path: Path) -> None:
    value = (
        "Проверку документов специалист производит в течение 24 часов с момента регистрации заявки. "
        "В случае если специалист должен проверить имущество перед согласованием, сотрудник "
        "обязан проставить отметку НЕ СОГЛАСОВАН в журнале и системе учета. "
        "В этом случае составляется акт несоответствия с расчетом суммы. "
        "С сотрудника берется объяснительная и обязательство, акты передаются в отдел контроля."
    )
    _facts, axes = await _converted_axes(tmp_path, value)
    assert any(a.correct_value == "в течение 24 часов с момента регистрации заявки" for a in axes)
    assert not any("НЕ СОГЛАСОВАН" in a.correct_value for a in axes)
