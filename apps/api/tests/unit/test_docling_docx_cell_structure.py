from __future__ import annotations

import importlib.util
import zipfile
from pathlib import Path

from docx import Document

SERVICE_PATH = Path(__file__).resolve().parents[4] / "infra" / "docling-service" / "main.py"


def _load_service(monkeypatch):
    monkeypatch.setenv("DOCLING_API_KEY", "test-docling-key")
    spec = importlib.util.spec_from_file_location("docling_docx_cell_structure_test", SERVICE_PATH)
    assert spec and spec.loader
    service = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(service)
    return service


def _make_docx(path: Path, *, with_table: bool) -> None:
    document = Document()
    document.add_paragraph("Text before the table.")
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Region A"
        table.cell(0, 0).add_paragraph("1. inspect limit")
        table.cell(0, 0).add_paragraph("- approve; record")
        table.cell(0, 1).text = "700|100 to 1|100|000"
        table.cell(0, 1).add_paragraph("Only after approval.")
        table.cell(1, 0).text = "Region B"
        table.cell(1, 0).add_paragraph("2. inspect outside city")
        table.cell(1, 1).text = "700 100 to 1 400 000"
        table.cell(1, 1).add_paragraph("If remote, call first.")
    document.add_paragraph("Text after the table.")
    document.save(path)


def _rewrite_document_xml(source: Path, destination: Path, transform) -> None:
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(destination, "w") as destination_zip:
        for item in source_zip.infolist():
            payload = source_zip.read(item.filename)
            if item.filename == "word/document.xml":
                payload = transform(payload)
            destination_zip.writestr(item, payload)


def _flattened_table_markdown(cell_left: str = "Region A", cell_right: str = "Only after approval.") -> str:
    return (
        "# Outside heading\n\n"
        "**Bold outside** [link](https://example.test)\\n\\\\path\n\n"
        f"| {cell_left} | {cell_right} |\n"
        "| --- | --- |\n"
        "| Region B | If remote, call first. |\n"
    )


def test_docx_table_conversion_preserves_cell_paragraph_boundaries_and_order(
    tmp_path: Path, monkeypatch
) -> None:
    service = _load_service(monkeypatch)
    source = tmp_path / "structured.docx"
    _make_docx(source, with_table=True)

    class FlatteningMarkItDown:
        def convert_local(self, _path):
            return type("Result", (), {"markdown": (
                "Text before the table.\n\n"
                "| Region A 1. inspect limit - approve; record | "
                "700\\|100 to 1\\|100\\|000 Only after approval. |\n"
                "| --- | --- |\n"
                "| Region B 2. inspect outside city | 700 100 to 1 400 000 If remote, call first. |\n\n"
                "Text after the table."
            )})()

    monkeypatch.setattr(service, "get_markitdown_converter", lambda: FlatteningMarkItDown())
    output = service._markitdown_convert(str(source))

    assert "Text before the table." in output
    assert "Text after the table." in output
    assert "Region A<br>1. inspect limit<br>- approve; record" in output
    assert "700\\|100 to 1\\|100\\|000<br>Only after approval." in output
    assert "Region B<br>2. inspect outside city" in output
    assert "If remote, call first." in output
    assert output.index("Text before the table.") < output.index("Region A") < output.index("Text after the table.")
    assert "Text before the table." in output


def test_docx_without_tables_retains_markitdown_output(
    tmp_path: Path, monkeypatch
) -> None:
    service = _load_service(monkeypatch)
    source = tmp_path / "plain.docx"
    _make_docx(source, with_table=False)

    class MarkItDown:
        def convert_local(self, _path):
            return type("Result", (), {"markdown": "MarkItDown's non-table output"})()

    monkeypatch.setattr(service, "get_markitdown_converter", lambda: MarkItDown())
    assert service._markitdown_convert(str(source)) == "MarkItDown's non-table output"


def test_docx_preserves_markitdown_formatting_outside_tables_and_generated_list_labels(
    tmp_path: Path, monkeypatch
) -> None:
    service = _load_service(monkeypatch)
    source = tmp_path / "lists-and-formatting.docx"
    document = Document()
    document.add_heading("Source heading", level=1)
    document.add_paragraph("Source bold and link wrapper")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Region A"
    numbered = table.cell(0, 0).add_paragraph("inspect first")
    numbered.style = "List Number"
    bulleted = table.cell(0, 0).add_paragraph("approve second")
    bulleted.style = "List Bullet"
    table.cell(0, 1).text = r"C:\docs\policy"
    document.save(source)

    class MarkItDown:
        def convert_local(self, _path):
            return type("Result", (), {"markdown": (
                "# Outside heading\n\n"
                "**Bold outside** [link](https://example.test)\\n\\\\path\n\n"
                "| Region A 1. inspect first - approve second | C:\\\\docs\\\\policy |\n"
                "| --- | --- |\n"
            )})()

    monkeypatch.setattr(service, "get_markitdown_converter", lambda: MarkItDown())
    output = service._markitdown_convert(str(source))

    assert "# Outside heading\n\n**Bold outside** [link](https://example.test)\\n\\\\path" in output
    assert "Region A<br>1. inspect first<br>- approve second" in output
    assert r"C:\\docs\\policy" in output


def test_docx_merged_cells_fail_closed_without_rewriting_markitdown(
    tmp_path: Path, monkeypatch
) -> None:
    service = _load_service(monkeypatch)
    source = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Merged source"
    table.cell(0, 0).merge(table.cell(0, 1))
    document.save(source)
    markdown = "| merged source | second cell |\n| --- | --- |\n"

    class MarkItDown:
        def convert_local(self, _path):
            return type("Result", (), {"markdown": markdown})()

    monkeypatch.setattr(service, "get_markitdown_converter", lambda: MarkItDown())
    assert service._markitdown_convert(str(source)) == markdown


def test_docx_content_control_wrapper_fails_closed_without_omitting_output(
    tmp_path: Path, monkeypatch
) -> None:
    service = _load_service(monkeypatch)
    plain = tmp_path / "plain.docx"
    wrapped = tmp_path / "content-control.docx"
    _make_docx(plain, with_table=True)

    def wrap_table(xml: bytes) -> bytes:
        start = xml.index(b"<w:tbl>")
        end = xml.index(b"</w:tbl>", start) + len(b"</w:tbl>")
        return xml[:start] + b"<w:sdt><w:sdtContent>" + xml[start:end] + b"</w:sdtContent></w:sdt>" + xml[end:]

    _rewrite_document_xml(plain, wrapped, wrap_table)
    markdown = "| a | b |\n| --- | --- |\n"

    class MarkItDown:
        def convert_local(self, _path):
            return type("Result", (), {"markdown": markdown})()

    monkeypatch.setattr(service, "get_markitdown_converter", lambda: MarkItDown())
    assert service._markitdown_convert(str(wrapped)) == markdown


def test_docx_entity_or_dtd_input_fails_closed_before_xml_expansion(
    tmp_path: Path, monkeypatch
) -> None:
    service = _load_service(monkeypatch)
    plain = tmp_path / "plain.docx"
    hostile = tmp_path / "entity.docx"
    _make_docx(plain, with_table=True)

    def add_dtd(xml: bytes) -> bytes:
        marker = b"?>"
        position = xml.index(marker) + len(marker)
        return xml[:position] + b'<!DOCTYPE w:document [<!ENTITY x "expanded">]>' + xml[position:]

    _rewrite_document_xml(plain, hostile, add_dtd)
    markdown = "| original | output |\n| --- | --- |\n"

    class MarkItDown:
        def convert_local(self, _path):
            return type("Result", (), {"markdown": markdown})()

    monkeypatch.setattr(service, "get_markitdown_converter", lambda: MarkItDown())
    assert service._markitdown_convert(str(hostile)) == markdown

    utf16_declared = tmp_path / "utf16-declared.docx"
    _rewrite_document_xml(
        plain,
        utf16_declared,
        lambda xml: b'<?xml version="1.0" encoding="utf-16"?>' + xml,
    )
    assert service._markitdown_convert(str(utf16_declared)) == markdown


def test_cell_matching_uses_exact_visible_characters_for_duplicates_breaks_and_formatted_numbers(
    monkeypatch,
) -> None:
    service = _load_service(monkeypatch)

    assert service._insert_cell_breaks(
        "same same", ["same", "same"]
    ) == "same<br>same"
    assert service._insert_cell_breaks(
        "before after", ["before<br>after"]
    ) == "before<br>after"
    assert service._insert_cell_breaks(
        "700 **100** to 1 400 000 .", ["700100 to 1400000", "."]
    ) == "700 **100** to 1 400 000<br>."
    assert service._insert_cell_breaks(
        "outsourced north ... wrong south", ["north", "south"]
    ) is None
    assert service._insert_cell_breaks("north", ["north", "south"]) is None
    assert service._insert_cell_breaks("north 42 south", ["north", "south"]) is None
