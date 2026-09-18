from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

from app.modules.ai.evidence_engine import CourseIntent, EvidenceCourseEngine
from app.modules.ai.evidence_engine.adapters import narrative_document_from_pages
from app.modules.ai.evidence_engine.engine import (
    _concise_narrative_title,
    _spreadsheet_question_prompt,
)
from app.modules.ai.evidence_engine.models import (
    AssessmentDraft,
    CourseDraft,
    LessonDraft,
    QuestionDraft,
    SourceDocument,
    SourceFact,
    SourceSection,
)
from app.modules.ai.evidence_engine.provider_engine import ProviderBackedEvidenceEngine
from app.modules.ai.evidence_engine.provider_models import ChatCompletion, EmbeddingBatch, GroundedBlock
from app.modules.ai.evidence_engine.quality import (
    evaluate_publishability,
    is_generic_question,
    question_has_ambiguous_options,
)


class _RecordingEmbeddings:
    def __init__(self) -> None:
        self.calls: list[list[str]] = []

    def embed(self, texts: list[str]) -> EmbeddingBatch:
        self.calls.append(list(texts))
        vectors = tuple(
            tuple(1.0 if index == position % 3 else 0.0 for index in range(3))
            for position, _ in enumerate(texts)
        )
        return EmbeddingBatch(vectors=vectors, model="qwen-test", duration_seconds=0.01)


class _UnavailableEmbeddings:
    def embed(self, texts: list[str]) -> EmbeddingBatch:
        del texts
        from app.modules.ai.evidence_engine.providers import ProviderCallError

        raise ProviderCallError("embedding endpoint unavailable")


class _GroundedChat:
    def __init__(self, *, unsupported_fact: bool = False) -> None:
        self.requests: list[dict[str, Any]] = []
        self.unsupported_fact = unsupported_fact

    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        self.requests.append(request)
        facts = request["facts"]
        fact_ids = [fact["fact_id"] for fact in facts]
        if self.unsupported_fact:
            fact_ids = ["fact-outside-plan"]
        questions = []
        for seed in request["question_seeds"]:
            questions.append(
                {
                    "prompt": f"Практическая проверка: {seed['prompt']}",
                    "options": seed["options"],
                    "correct_answer": seed["correct_answer"],
                    "explanation": "Ответ следует из указанного положения.",
                    "fact_ids": [seed["fact_id"]],
                }
            )
        payload = {
            "title": request["lesson_title"],
            "objective": request["objective"],
            "blocks": [
                {
                    "heading": "Основное правило",
                    "text": " ".join(fact["value"] for fact in facts),
                    "fact_ids": fact_ids,
                }
            ],
            "questions": questions,
        }
        return ChatCompletion(
            payload=payload,
            model="deepseek-test",
            duration_seconds=0.02,
            prompt_tokens=100,
            completion_tokens=80,
        )


class _MetadataNumberChat:
    def __init__(self, *, hallucinate: bool = False) -> None:
        self.hallucinate = hallucinate

    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        fact_id = request["facts"][0]["fact_id"]
        text = (
            "В разделе 4 дайте покупателю ответ в 2–3 предложениях."
            if not self.hallucinate
            else "В разделе 4 ответ нужно предоставить через 99 дней."
        )
        return ChatCompletion(
            payload={
                "title": request["lesson_title"],
                "objective": request["objective"],
                "blocks": [
                    {
                        "heading": "Применение",
                        "text": text,
                        "fact_ids": [fact_id],
                    }
                ],
                "questions": [],
            },
            model="deepseek-test",
            duration_seconds=0.01,
        )


class _NormalizedOcrNumberChat:
    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        fact_ids = [fact["fact_id"] for fact in request["facts"]]
        source_text = " ".join(fact["value"] for fact in request["facts"])
        normalized_parts = []
        if "трех" in source_text:
            normalized_parts.append("Уведомление направляется за 3 рабочих дня.")
        if "три десятых" in source_text:
            normalized_parts.append("Ставка составляет 0.3 процента.")
        return ChatCompletion(
            payload={
                "title": request["lesson_title"],
                "objective": request["objective"],
                "blocks": [
                    {
                        "heading": "Срок и ставка",
                        "text": " ".join(normalized_parts),
                        "fact_ids": fact_ids,
                    }
                ],
                "questions": [],
            },
            model="deepseek-test",
            duration_seconds=0.01,
        )


class _AutonomousOcrRepairChat:
    def __init__(self) -> None:
        self.attempts = 0

    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        self.attempts += 1
        fact_ids = [fact["fact_id"] for fact in request["facts"]]
        if self.attempts == 1:
            title = "После получения заявления сотрудник ломбарда обязан"
            text = "Порог ()()() МРП применяется при принятии решения."
        else:
            title = "Условия рассмотрения заявления"
            text = (
                "Сотрудник применяет только читаемые условия положения. "
                "Нераспознанное пороговое значение не воспроизводится и не заменяется догадкой."
            )
        return ChatCompletion(
            payload={
                "title": title,
                "objective": request["objective"],
                "blocks": [
                    {
                        "heading": "Применение правила",
                        "text": text,
                        "fact_ids": fact_ids,
                    }
                ],
                "questions": [],
            },
            model="deepseek-test",
            duration_seconds=0.01,
        )


class _InstructionLeakRepairChat(_GroundedChat):
    def __init__(self) -> None:
        super().__init__()
        self.attempts = 0

    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        self.attempts += 1
        completion = super().complete_json(request)
        if self.attempts == 1:
            completion.payload["blocks"][0]["text"] = (
                "Неверные варианты должны быть правдоподобными. "
                "Вопрос не должен содержать подсказок."
            )
        return completion


class _MetaQuestionChat(_GroundedChat):
    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        completion = super().complete_json(request)
        for question in completion.payload["questions"]:
            question["prompt"] = "О чём этот урок?"
        return completion


class _MutatingAnswerChat(_GroundedChat):
    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        completion = super().complete_json(request)
        for question in completion.payload["questions"]:
            question["options"] = ["Изменённый моделью вариант"]
            question["correct_answer"] = "Изменённый моделью ответ"
        return completion


class _OmittingQuestionRewriteChat(_GroundedChat):
    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        completion = super().complete_json(request)
        completion.payload["questions"] = []
        return completion


class _ExtraQuestionRewriteChat(_GroundedChat):
    def complete_json(self, request: dict[str, Any]) -> ChatCompletion:
        completion = super().complete_json(request)
        completion.payload["questions"].append(
            {
                "prompt": "Лишний вопрос модели",
                "explanation": "Не относится к серверному плану.",
                "fact_ids": ["unknown-fact"],
            }
        )
        return completion


def _narrative_source(*facts: str) -> SourceDocument:
    return SourceDocument(
        source_id="policy",
        title="Правила обслуживания клиентов",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="policy:section:1",
                title="Рассмотрение заявлений",
                role="primary",
                facts=tuple(
                    SourceFact(
                        fact_id=f"policy:fact:{index}",
                        subject="Рассмотрение заявлений",
                        attribute="правило",
                        value=value,
                        source_locator=f"page=1;sentence={index}",
                    )
                    for index, value in enumerate(facts, start=1)
                ),
            ),
        ),
    )


def test_small_source_is_not_padded_to_a_fixed_lesson_or_question_quota() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение пятнадцати рабочих дней.",
        "Сотрудник проверяет полноту представленных документов.",
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert len(result.course.lessons) == 1
    assert len(result.assessment.questions) <= 2
    assert result.evaluation.quota_padding_count == 0


def test_lesson_content_does_not_repeat_the_outer_lesson_heading() -> None:
    source = _narrative_source(
        "Сотрудник проверяет полноту представленных документов.",
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert len(result.course.lessons) == 1
    assert not result.course.lessons[0].content.startswith("## ")
    assert result.course.lessons[0].content.startswith("### Подтверждённые сведения")


def test_lessons_and_questions_share_the_same_admitted_fact_ids() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Уведомление направляется не позднее 30 календарных дней.",
        "Ответ на короткий запрос готовится в течение 3 рабочих дней.",
    )

    result = EvidenceCourseEngine().generate_from_document(source)
    lesson_fact_ids = {
        fact_id for lesson in result.course.lessons for fact_id in lesson.fact_ids
    }

    assert lesson_fact_ids
    assert all(question.fact_id in lesson_fact_ids for question in result.assessment.questions)
    assert all(
        question.lesson_id
        == next(
            lesson.lesson_id
            for lesson in result.course.lessons
            if question.fact_id in lesson.fact_ids
        )
        for question in result.assessment.questions
    )


def test_question_deduplication_is_fact_based_and_order_independent() -> None:
    duplicate = "Жалоба рассматривается в течение тридцати календарных дней."
    first = _narrative_source(duplicate, duplicate, "Ответ направляется клиенту письменно.")
    second = _narrative_source("Ответ направляется клиенту письменно.", duplicate, duplicate)

    first_result = EvidenceCourseEngine().generate_from_document(first)
    second_result = EvidenceCourseEngine().generate_from_document(second)

    assert first_result.semantic_fingerprint == second_result.semantic_fingerprint
    assert first_result.evaluation.duplicate_fact_count == 1
    assert second_result.evaluation.duplicate_fact_count == 1


def test_repeated_runs_have_stable_semantic_fingerprint() -> None:
    source = _narrative_source(
        "Клиент вправе получить полную информацию об условиях микрокредита.",
        "Ломбард обеспечивает конфиденциальность сведений о клиенте.",
        "Обращение регистрируется в день поступления.",
    )

    fingerprints = {
        EvidenceCourseEngine().generate_from_document(source, simulation_seed=seed).semantic_fingerprint
        for seed in range(5)
    }

    assert len(fingerprints) == 1


def test_spreadsheet_plan_respects_source_derived_teachable_capacity() -> None:
    source = SourceDocument(
        source_id="bounded-catalog",
        title="Каталог материалов",
        kind="spreadsheet",
        teachable_units=3,
        sections=(
            SourceSection(
                section_id="materials",
                title="Материалы",
                role="primary",
                facts=tuple(
                    SourceFact(
                        fact_id=f"material-{index}",
                        subject=f"Материал {index}",
                        attribute="Применение",
                        value=f"Материал {index} применяется для подтвержденной операции {index}.",
                        source_locator=f"sheet=Материалы;row={index + 1}",
                    )
                    for index in range(1, 9)
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert 1 <= len(result.evidence_plan) <= source.teachable_units
    assert {
        fact_id
        for lesson in result.evidence_plan
        for fact_id in lesson.fact_ids
    } == {fact.fact_id for fact in result.admitted_facts}


def test_spreadsheet_plan_bounds_multi_module_capacity_without_repeating_facts() -> None:
    sections = tuple(
        SourceSection(
            section_id=f"module-{module_index}",
            title=f"Модуль {module_index}",
            role="primary",
            facts=tuple(
                SourceFact(
                    fact_id=f"fact-{module_index}-{subject_index}",
                    subject=f"Тема {module_index}-{subject_index}",
                    attribute="Применение",
                    value=(
                        f"Тема {module_index}-{subject_index} применяется для "
                        f"подтвержденной операции {module_index}-{subject_index}."
                    ),
                    source_locator=f"sheet=Модуль {module_index};row={subject_index}",
                )
                for subject_index in range(1, 5)
            ),
        )
        for module_index in range(1, 3)
    )
    source = SourceDocument(
        source_id="bounded-multi-module-catalog",
        title="Каталог материалов",
        kind="spreadsheet",
        teachable_units=3,
        sections=sections,
    )

    result = EvidenceCourseEngine().generate_from_document(source)
    flattened_fact_ids = tuple(
        fact_id for lesson in result.evidence_plan for fact_id in lesson.fact_ids
    )

    assert len(result.evidence_plan) <= source.teachable_units
    assert flattened_fact_ids == tuple(fact.fact_id for fact in result.admitted_facts)
    assert len(flattened_fact_ids) == len(set(flattened_fact_ids))
    assert {lesson.module_title for lesson in result.evidence_plan} <= {
        "Модуль 1",
        "Модуль 2",
    }
    assert len({lesson.title for lesson in result.evidence_plan}) == len(
        result.evidence_plan
    )


def test_narrative_plan_deduplicates_mapped_fact_ids_in_first_source_order() -> None:
    shared_value = "Заявление рассматривается в течение 15 рабочих дней."
    source = SourceDocument(
        source_id="repeated-narrative-facts",
        title="Порядок рассмотрения",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="section-one",
                title="1. Первый раздел",
                role="primary",
                facts=(
                    SourceFact(
                        fact_id="raw-first",
                        subject="Заявление",
                        attribute="срок рассмотрения",
                        value=shared_value,
                        source_locator="page=1;line=1",
                    ),
                    SourceFact(
                        fact_id="raw-second",
                        subject="Ответ",
                        attribute="срок ответа",
                        value="Ответ направляется в течение 30 календарных дней.",
                        source_locator="page=1;line=2",
                    ),
                    SourceFact(
                        fact_id="raw-first-duplicate",
                        subject="Заявление",
                        attribute="срок рассмотрения",
                        value=shared_value,
                        source_locator="page=1;line=3",
                    ),
                ),
            ),
            SourceSection(
                section_id="section-two",
                title="2. Второй раздел",
                role="primary",
                facts=(
                    SourceFact(
                        fact_id="raw-second-duplicate",
                        subject="Ответ",
                        attribute="срок ответа",
                        value="Ответ направляется в течение 30 календарных дней.",
                        source_locator="page=2;line=1",
                    ),
                    SourceFact(
                        fact_id="raw-third",
                        subject="Уведомление",
                        attribute="срок уведомления",
                        value="Уведомление направляется в течение 3 рабочих дней.",
                        source_locator="page=2;line=2",
                    ),
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)
    flattened_fact_ids = tuple(
        fact_id for lesson in result.evidence_plan for fact_id in lesson.fact_ids
    )

    assert flattened_fact_ids == tuple(fact.fact_id for fact in result.admitted_facts)
    assert len(flattened_fact_ids) == len(set(flattened_fact_ids))


def test_large_narrative_preserves_source_sections_without_partition_explosion() -> None:
    sections = tuple(
        SourceSection(
            section_id=f"section-{section_index}",
            title=f"{section_index}. Раздел {section_index}",
            role="primary",
            facts=tuple(
                SourceFact(
                    fact_id=f"fact-{section_index}-{fact_index}",
                    subject=f"Раздел {section_index}",
                    attribute="правило",
                    value=(
                        f"Сотрудник выполняет подтвержденное действие {fact_index} "
                        f"в разделе {section_index}."
                    ),
                    source_locator=f"section={section_index};fact={fact_index}",
                )
                for fact_index in range(1, 37 if section_index == 17 else 9)
            ),
        )
        for section_index in range(1, 18)
    )
    source = SourceDocument(
        source_id="large-policy",
        title="Правила склада",
        kind="narrative",
        sections=sections,
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert 9 <= len(result.evidence_plan) <= 12
    assert max(len(lesson.fact_ids) for lesson in result.evidence_plan) <= 20
    assert {
        fact_id
        for lesson in result.evidence_plan
        for fact_id in lesson.fact_ids
    } == {fact.fact_id for fact in result.admitted_facts}


def test_narrative_plan_hard_bounds_source_words_for_safe_fallback() -> None:
    repeated = " ".join(f"слово{index}" for index in range(70))
    source = SourceDocument(
        source_id="fallback-word-bound",
        title="Большой раздел",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="large-section",
                title="1. БОЛЬШОЙ РАЗДЕЛ",
                role="primary",
                facts=tuple(
                    SourceFact(
                        fact_id=f"fact-{index}",
                        subject="Большой раздел",
                        attribute="правило",
                        value=f"{repeated} пункт {index}",
                        source_locator=f"section=1;clause={index}",
                    )
                    for index in range(1, 13)
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)
    facts_by_id = {fact.fact_id: fact for fact in result.admitted_facts}
    lesson_source_words = [
        sum(len(facts_by_id[fact_id].value.split()) for fact_id in lesson.fact_ids)
        for lesson in result.evidence_plan
    ]

    assert max(lesson_source_words) <= 550
    assert {
        fact_id
        for lesson in result.evidence_plan
        for fact_id in lesson.fact_ids
    } == set(facts_by_id)


def test_long_policy_uses_coherent_lesson_density_without_losing_rules() -> None:
    facts_per_section = (14, 23, 20, 7, 13, 26, 12, 2, 22, 71, 10, 6)
    sections = tuple(
        SourceSection(
            section_id=f"policy-section-{section_index}",
            title=f"{section_index}. Раздел правил {section_index}",
            role="primary",
            facts=tuple(
                SourceFact(
                    fact_id=f"policy-fact-{section_index}-{fact_index}",
                    subject=f"Раздел правил {section_index}",
                    attribute="правило",
                    value=(
                        "Сотрудник проверяет заявление, подтверждающие документы и условия "
                        "операции, фиксирует принятое решение в учётной системе, уведомляет "
                        "клиента в установленном порядке и сохраняет подтверждение выполненного "
                        f"действия по пункту {section_index}.{fact_index}."
                    ),
                    source_locator=f"section={section_index};clause={fact_index}",
                )
                for fact_index in range(1, fact_count + 1)
            ),
        )
        for section_index, fact_count in enumerate(facts_per_section, start=1)
    )
    source = SourceDocument(
        source_id="full-policy",
        title="Правила предоставления услуг",
        kind="narrative",
        sections=sections,
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert len(sections) <= len(result.evidence_plan) <= 20
    assert max(len(lesson.fact_ids) for lesson in result.evidence_plan) <= 20
    assert len({lesson.title for lesson in result.evidence_plan}) == len(
        result.evidence_plan
    )
    assert {
        fact_id
        for lesson in result.evidence_plan
        for fact_id in lesson.fact_ids
    } == {fact.fact_id for fact in result.admitted_facts}


def test_split_narrative_section_has_unique_part_titles() -> None:
    repeated_text = " ".join(
        (
            "Сотрудник",
            "проверяет",
            "документы",
            "фиксирует",
            "решение",
            "уведомляет",
            "клиента",
            "сохраняет",
            "подтверждение",
            "контролирует",
        )
        * 4
    )
    source = SourceDocument(
        source_id="large-section",
        title="Права и обязанности",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="rights",
                title="10. Права и обязанности сторон",
                role="primary",
                facts=tuple(
                    SourceFact(
                        fact_id=f"right-{index}",
                        subject="Права и обязанности сторон",
                        attribute="правило",
                        value=f"{repeated_text} Пункт {index}.",
                        source_locator=f"clause={index}",
                    )
                    for index in range(1, 43)
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)
    titles = [lesson.title for lesson in result.evidence_plan]

    assert len(titles) > 1
    assert len(set(titles)) == len(titles)


def test_legal_section_titles_are_concise_learning_topics() -> None:
    cases = {
        "2. ПОРЯДОК ПОДАЧИ ЗАЯВЛЕНИЯ НА ПРЕДОСТАВЛЕНИЕ МИКРОКРЕДИТА И ПОРЯДОК ЕГО РАССМОТРЕНИЯ": (
            "Подача и рассмотрение заявления"
        ),
        "7. ПРАВИЛА РАСЧЕТА ГОДОВОЙ ЭФФЕКТИВНОЙ СТАВКИ ВОЗНАГРАЖДЕНИЯ ПО ПРЕДОСТАВЛЯЕМЫМ МИКРОКРЕДИТАМ": (
            "Расчёт годовой эффективной ставки"
        ),
        "10. ПРАВА И ОБЯЗАННОСТИ ЛОМБАРДА И ЕГО КЛИЕНТОВ, ИХ ВЗАИМНАЯ ОТВЕТСТВЕННОСТЬ": (
            "Права, обязанности и ответственность сторон"
        ),
    }

    assert {
        source: _concise_narrative_title(source) for source in cases
    } == cases


def test_xlsx_primary_matrix_drives_curriculum_and_catalog_is_supporting(tmp_path: Path) -> None:
    workbook = Workbook()
    primary = workbook.active
    primary.title = "Коллекции"
    primary.append(["Поле", "Феникс", "Чикаго Нео", "Чикаго Стрит"])
    primary.append(["Назначение", "Для хранения", "Для спальни", "Для прихожей"])
    primary.append(["Материал", "ЛДСП", "ЛДСП и металл", "ЛДСП"])
    primary.append(["Цвета", "Белый и графит", "Кашемир", "Бетон"])
    catalog = workbook.create_sheet("Феникс и Чикаго")
    catalog.append(["Артикул", "Наименование", "Вес", "Описание"])
    for index in range(80):
        catalog.append(
            [f"SKU-{index:03d}", f"Система Феникс модуль {index}", index + 1, "Карточка товара"]
        )
    source_path = tmp_path / "catalog.xlsx"
    workbook.save(source_path)

    result = EvidenceCourseEngine().generate(source_path, intent=CourseIntent())

    lesson_titles = {lesson.title for lesson in result.course.lessons}
    assert any("Феникс" in title for title in lesson_titles)
    assert any("Чикаго Нео" in title for title in lesson_titles)
    assert any("Чикаго Стрит" in title for title in lesson_titles)
    assert not any("SKU" in title or "Артикул" in title for title in lesson_titles)
    assert result.document_plan.primary_sections == ("Коллекции",)
    assert result.document_plan.supporting_sections == ("Феникс и Чикаго",)
    assert result.document_plan.supporting_fact_count > 0
    assert any(plan.supporting_fact_ids for plan in result.evidence_plan)


def test_safe_assessment_omits_question_when_compatible_distractors_do_not_exist() -> None:
    source = _narrative_source(
        "Ломбард обеспечивает конфиденциальность сведений о клиенте.",
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert len(result.course.lessons) == 1
    assert result.assessment.questions == ()
    assert result.evaluation.warnings == ("assessment_empty_no_safe_questions",)
    assert result.evaluation.unsupported_claim_count == 0


def test_narrative_non_numeric_rules_are_not_padded_with_generic_questions() -> None:
    source = SourceDocument(
        source_id="policy-rules",
        title="Правила",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="rights",
                title="10. Права и обязанности",
                role="primary",
                facts=tuple(
                    SourceFact(
                        fact_id=f"right-{index}",
                        subject="10. Права и обязанности",
                        attribute="обязанность",
                        value=value,
                        source_locator=f"section=10;clause={index}",
                    )
                    for index, value in enumerate(
                        (
                            "Ломбард обязан выдать клиенту подписанный залоговый билет.",
                            "Ломбард обязан обеспечить сохранность принятого залога.",
                            "Ломбард обязан предоставить клиенту полную информацию об условиях.",
                        ),
                        start=1,
                    )
                ),
            ),
            SourceSection(
                section_id="privacy",
                title="12. Конфиденциальность",
                role="primary",
                facts=(
                    SourceFact(
                        fact_id="privacy-1",
                        subject="12. Конфиденциальность",
                        attribute="обязанность",
                        value="Ломбард обязан защищать персональные данные клиента.",
                        source_locator="section=12;clause=1",
                    ),
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)
    assert result.assessment.questions == ()


def test_narrative_questions_mask_exact_values_and_use_only_sourced_distractors() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Уведомление направляется не позднее 30 календарных дней.",
        "Ответ на короткий запрос готовится в течение 3 рабочих дней.",
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert len(result.assessment.questions) == 3
    assert len({question.prompt for question in result.assessment.questions}) == 3
    assert all("_____" in question.prompt for question in result.assessment.questions)
    sourced_values = {"15 рабочих дней", "30 календарных дней", "3 рабочих дней"}
    assert all(set(question.options) == sourced_values for question in result.assessment.questions)


def test_narrative_duration_distractors_do_not_mix_minutes_with_days() -> None:
    source = _narrative_source(
        "Первая проверка выполняется в течение 15 минут.",
        "Вторая проверка выполняется в течение 20 минут.",
        "Третья проверка выполняется в течение 45 минут.",
        "Первый ответ направляется в течение 3 рабочих дней.",
        "Второй ответ направляется в течение 7 дней.",
        "Третий ответ направляется в течение 30 календарных дней.",
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    for question in result.assessment.questions:
        rendered = " ".join(question.options).casefold()
        has_short = "минут" in rendered or "час" in rendered
        has_calendar = "дн" in rendered or "месяц" in rendered
        assert not (has_short and has_calendar)


def test_narrative_options_do_not_include_equivalent_numeric_spellings() -> None:
    source = _narrative_source(
        "Первое уведомление направляется в течение 30 (тридцати) календарных дней.",
        "Второе уведомление направляется в течение тридцати календарных дней.",
        "Ответ направляется в течение 15 рабочих дней.",
        "Короткое сообщение направляется в течение семи дней.",
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    for question in result.assessment.questions:
        rendered = " | ".join(question.options).casefold()
        assert not (
            "30 (тридцати) календарных дней" in rendered
            and "тридцати календарных дней" in rendered
        )


def test_narrative_adapter_joins_wrapped_clauses_and_keeps_major_section_scope(
    tmp_path: Path,
) -> None:
    source_path = tmp_path / "rules.pdf"
    source_path.write_bytes(b"immutable-test-source")
    document = narrative_document_from_pages(
        path=source_path,
        pages=[
            "1. ОБЩИЕ ПОЛОЖЕНИЯ\n"
            "1. Настоящие правила определяют порядок предоставления\n"
            "микрокредитов физическим лицам под залог имущества.\n"
            "2. Правила предоставляются клиенту по первому требованию.",
            "2. ПОРЯДОК ПОДАЧИ ЗАЯВЛЕНИЯ\n"
            "1. Заявление рассматривается в течение пятнадцати\n"
            "рабочих дней со дня регистрации.",
            "Приложение № 1\n"
            "Заявление на предоставление микрокредита\n"
            "ФИО ________________________",
        ],
    )

    assert [section.title for section in document.sections if section.role == "primary"] == [
        "1. ОБЩИЕ ПОЛОЖЕНИЯ",
        "2. ПОРЯДОК ПОДАЧИ ЗАЯВЛЕНИЯ",
    ]
    first_fact = document.sections[0].facts[0]
    assert "порядок предоставления микрокредитов" in first_fact.value
    assert "под залог имущества" in first_fact.value
    assert first_fact.source_locator.startswith("page=1")
    assert any(section.role == "supporting" for section in document.sections)


def test_provider_v2_indexes_documents_as_is_and_prefixes_only_queries() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )
    embeddings = _RecordingEmbeddings()
    chat = _GroundedChat()

    result = ProviderBackedEvidenceEngine(
        embeddings=embeddings,
        chat=chat,
    ).generate_from_document(source)

    assert embeddings.calls[0] == [fact.value for fact in source.sections[0].facts]
    assert all(
        text.startswith(
            "Instruct: Given a user question, retrieve relevant passages that answer the question\nQuery: "
        )
        for text in embeddings.calls[1]
    )
    assert result.embedding_dimension == 3
    assert result.provider_fallback_count == 0
    assert len(result.realized_course.lessons) == 1


def test_provider_v2_rejects_fact_ids_outside_lesson_and_keeps_draft_fallback() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )
    embeddings = _RecordingEmbeddings()
    chat = _GroundedChat(unsupported_fact=True)

    result = ProviderBackedEvidenceEngine(
        embeddings=embeddings,
        chat=chat,
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 1
    assert result.realized_course.lessons == result.evidence_result.course.lessons
    assert result.validation_errors[0].startswith("lesson-")


def test_provider_v2_preserves_source_backed_answer_options() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )
    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_GroundedChat(),
    ).generate_from_document(source)

    original = result.evidence_result.assessment.questions
    realized = result.realized_assessment.questions
    assert len(realized) == len(original)
    assert all(
        set(new.options) == set(seed.options)
        and new.correct_answer == seed.correct_answer
        and new.fact_id == seed.fact_id
        for new, seed in zip(realized, original, strict=True)
    )


def test_provider_v2_treats_answer_options_as_server_owned_data() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_MutatingAnswerChat(),
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 0
    assert result.publishability.publishable is True
    assert all(
        realized.options == seed.options and realized.correct_answer == seed.correct_answer
        for realized, seed in zip(
            result.realized_assessment.questions,
            result.evidence_result.assessment.questions,
            strict=True,
        )
    )


def test_provider_v2_keeps_safe_seed_when_model_omits_optional_question_rewrite() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_OmittingQuestionRewriteChat(),
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 0
    assert result.publishability.publishable is True
    assert result.realized_assessment.questions == result.evidence_result.assessment.questions


def test_provider_v2_allows_numbers_present_in_fact_subject_or_attribute() -> None:
    source = SourceDocument(
        source_id="metadata-numbers",
        title="Ответ покупателю",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="section-4",
                title="Раздел 4",
                role="primary",
                facts=(
                    SourceFact(
                        fact_id="raw-fact",
                        subject="Раздел 4",
                        attribute="Ответ в 2–3 предложениях",
                        value="Сообщите покупателю подтверждённые преимущества коллекции.",
                        source_locator="page=4",
                    ),
                ),
            ),
        ),
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_MetadataNumberChat(),
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 0


def test_provider_v2_rejects_number_absent_from_all_cited_source_metadata() -> None:
    source = SourceDocument(
        source_id="hallucinated-number",
        title="Ответ покупателю",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="section-4",
                title="Раздел 4",
                role="primary",
                facts=(
                    SourceFact(
                        fact_id="raw-fact",
                        subject="Раздел 4",
                        attribute="Ответ в 2–3 предложениях",
                        value="Сообщите покупателю подтверждённые преимущества коллекции.",
                        source_locator="page=4",
                    ),
                ),
            ),
        ),
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_MetadataNumberChat(hallucinate=True),
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 1
    assert "introduces a number" in result.validation_errors[0]


def test_provider_v2_accepts_digit_normalization_of_russian_and_ocr_numbers() -> None:
    source = SourceDocument(
        source_id="ocr-number-normalization",
        title="Срок и ставка",
        kind="narrative",
        sections=(
            SourceSection(
                section_id="section",
                title="Условия",
                role="primary",
                facts=(
                    SourceFact(
                        fact_id="raw-one",
                        subject="Уведомление",
                        attribute="срок",
                        value="Клиент извещается в течение З (трех) рабочих дней.",
                        source_locator="page=1",
                    ),
                    SourceFact(
                        fact_id="raw-two",
                        subject="Ставка",
                        attribute="финансовое условие",
                        value="Ставка составляет (),З (ноль целых три десятых) процента.",
                        source_locator="page=1",
                    ),
                ),
            ),
        ),
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_NormalizedOcrNumberChat(),
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 0


def test_provider_v2_returns_publishable_grounded_blocks_for_clean_result() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_GroundedChat(),
    ).generate_from_document(source)

    assert result.publishability.publishable is True
    assert result.publishability.reasons == ()
    assert result.publishability.fact_coverage_ratio == 1.0
    assert result.grounded_blocks
    assert {fact_id for block in result.grounded_blocks for fact_id in block.fact_ids} == {
        fact.fact_id for fact in result.evidence_result.admitted_facts
    }


def test_provider_v2_repairs_ocr_noise_without_methodologist_preflight() -> None:
    source = _narrative_source(
        "Микрокредит не предоставляется выше ()()() месячного расчётного показателя.",
        "Заявление рассматривается сотрудником ломбарда.",
    )
    chat = _AutonomousOcrRepairChat()

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=chat,
        max_realizer_attempts=2,
    ).generate_from_document(source)

    assert chat.attempts == 2
    assert result.provider_fallback_count == 0
    assert result.publishability.publishable is True
    assert result.publishability.ocr_artifact_count == 0
    assert "()()()" not in result.realized_course.lessons[0].content
    assert result.realized_course.lessons[0].title == "Условия рассмотрения заявления"


def test_provider_v2_drops_meta_question_rewrite_and_keeps_safe_seed() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_MetaQuestionChat(),
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 0
    assert result.publishability.publishable is True
    assert all("О чём этот урок" not in question.prompt for question in result.realized_assessment.questions)


def test_customer_reported_source_meta_question_is_generic() -> None:
    assert is_generic_question(
        "Что в материале урока указано о выдаче микрокредита после проверки заявления?"
    )


def test_provider_v2_retries_internal_generation_instruction_leak() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )
    chat = _InstructionLeakRepairChat()

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=chat,
        max_realizer_attempts=2,
    ).generate_from_document(source)

    assert chat.attempts == 2
    assert result.provider_fallback_count == 0
    content = result.realized_course.lessons[0].content.casefold()
    assert "неверные варианты" not in content
    assert "вопрос не должен содержать подсказок" not in content


def test_provider_v2_ignores_extra_model_questions_outside_server_plan() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_ExtraQuestionRewriteChat(),
        max_realizer_attempts=1,
    ).generate_from_document(source)

    assert result.provider_fallback_count == 0
    assert result.publishability.publishable is True
    assert len(result.realized_assessment.questions) == len(result.evidence_result.assessment.questions)
    assert all(question.fact_id != "unknown-fact" for question in result.realized_assessment.questions)


def test_provider_v2_uses_subject_specific_spreadsheet_questions() -> None:
    source = SourceDocument(
        source_id="collections",
        title="Коллекции",
        kind="spreadsheet",
        sections=(
            SourceSection(
                section_id="primary",
                title="Коллекции",
                role="primary",
                facts=(
                    SourceFact("one", "Феникс", "Для каких комнат", "Спальня и гостиная", "row=1"),
                    SourceFact("two", "Чикаго", "Для каких комнат", "Прихожая и спальня", "row=2"),
                    SourceFact("three", "Нео", "Для каких комнат", "Гостиная и кабинет", "row=3"),
                ),
            ),
        ),
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_RecordingEmbeddings(),
        chat=_GroundedChat(),
    ).generate_from_document(source)

    prompts = [question.prompt for question in result.realized_assessment.questions]
    assert prompts
    assert all("Что указано" not in prompt for prompt in prompts)
    assert any("Для каких помещений подходит" in prompt for prompt in prompts)


def test_spreadsheet_question_prompts_are_natural_for_operational_attributes() -> None:
    cases = (
        ("Ширина", "АЛ-НШ-300", "Какова ширина «АЛ-НШ-300»?"),
        ("Где используется", "ПЭТ-плёнка", "Где используется «ПЭТ-плёнка»?"),
        (
            "Решение",
            "Клиенту нужна узкая прихожая",
            "Какое решение следует применить в ситуации «Клиенту нужна узкая прихожая»?",
        ),
        (
            "Почему",
            "Клиенту нужна узкая прихожая",
            "Почему для ситуации «Клиенту нужна узкая прихожая» рекомендуется указанное решение?",
        ),
        (
            "Особенности шкафов",
            "Чикаго Нео",
            "Какие особенности шкафов характерны для «Чикаго Нео»?",
        ),
        (
            "Главное преимущество",
            "Чикаго Стрит",
            "Какое главное преимущество лучше всего характеризует «Чикаго Стрит»?",
        ),
    )

    for index, (attribute, subject, expected) in enumerate(cases):
        fact = SourceFact(
            fact_id=f"fact-{index}",
            subject=subject,
            attribute=attribute,
            value="Проверяемое значение",
            source_locator=f"row={index + 1}",
        )
        assert _spreadsheet_question_prompt(fact) == expected


def test_spreadsheet_question_prompt_has_a_readable_specific_fallback() -> None:
    fact = SourceFact(
        fact_id="fact-custom",
        subject="Модуль М-1",
        attribute="Нагрузка на полку",
        value="До 10 кг",
        source_locator="row=1",
    )

    assert _spreadsheet_question_prompt(fact) == (
        "Какое значение параметра «Нагрузка на полку» установлено для «Модуль М-1»?"
    )


def test_spreadsheet_assessment_drops_answer_exposed_by_subject_identifier() -> None:
    source = SourceDocument(
        source_id="encoded-sku-widths",
        title="Размеры модулей",
        kind="spreadsheet",
        sections=(
            SourceSection(
                section_id="modules",
                title="Модули",
                role="primary",
                facts=(
                    SourceFact("width-300", "АЛ-НШ-300", "Ширина", "300", "row=1"),
                    SourceFact("width-400", "АЛ-ВШ-400", "Ширина", "400", "row=2"),
                    SourceFact("width-500", "АЛ-ЯЩ-500", "Ширина", "500", "row=3"),
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert result.assessment.questions == ()


def test_spreadsheet_assessment_drops_catalog_identifier_recall() -> None:
    source = SourceDocument(
        source_id="catalog-identifiers",
        title="Каталог модулей",
        kind="spreadsheet",
        sections=(
            SourceSection(
                section_id="modules",
                title="Модули",
                role="primary",
                facts=(
                    SourceFact("collection-a", "Модуль A", "Коллекция", "Альта", "row=1"),
                    SourceFact("collection-b", "Модуль B", "Коллекция", "Нордик", "row=2"),
                    SourceFact("collection-c", "Модуль C", "Коллекция", "Терра", "row=3"),
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    assert result.assessment.questions == ()


def test_spreadsheet_assessment_does_not_repeat_same_attribute_answer() -> None:
    source = SourceDocument(
        source_id="repeated-facades",
        title="Каталог фасадов",
        kind="spreadsheet",
        sections=(
            SourceSection(
                section_id="modules",
                title="Модули",
                role="primary",
                facts=(
                    SourceFact("facade-a", "Модуль A", "Фасад", "МДФ", "row=1"),
                    SourceFact("facade-b", "Модуль B", "Фасад", "МДФ", "row=2"),
                    SourceFact("facade-c", "Модуль C", "Фасад", "МДФ", "row=3"),
                    SourceFact("facade-d", "Модуль D", "Фасад", "Стекло", "row=4"),
                    SourceFact("facade-e", "Модуль E", "Фасад", "Шпон", "row=5"),
                ),
            ),
        ),
    )

    result = EvidenceCourseEngine().generate_from_document(source)

    keyed_answers = [
        (
            " ".join(question.prompt.casefold().split()),
            " ".join(question.correct_answer.casefold().split()),
        )
        for question in result.assessment.questions
    ]
    assert len(keyed_answers) == len(set(keyed_answers))
    assert [answer for _prompt, answer in keyed_answers].count("мдф") == 1


def test_provider_v2_continues_when_auxiliary_embedding_is_unavailable() -> None:
    source = _narrative_source(
        "Заявление рассматривается в течение 15 рабочих дней.",
        "Ответ направляется в течение 30 календарных дней.",
        "Короткое уведомление направляется в течение 3 рабочих дней.",
    )

    result = ProviderBackedEvidenceEngine(
        embeddings=_UnavailableEmbeddings(),
        chat=_GroundedChat(),
    ).generate_from_document(source)

    assert result.embedding_degraded is True
    assert result.embedding_error == "ProviderCallError"
    assert result.embedding_dimension == 0
    assert result.retrieval == ()
    assert result.provider_fallback_count == 0
    assert result.publishability.publishable is True


def test_v2_final_publishability_blocks_captured_style_and_ambiguous_answer_defects() -> None:
    lesson = LessonDraft(
        lesson_id="lesson-guides",
        module_title="Коллекции",
        title="Чикаго Стрит: направляющие",
        objective="Объяснять устройство направляющих",
        content="Полноценные габариты, а не компакт с маркетплейса.",
        fact_ids=("fact-guides",),
        supporting_fact_ids=(),
        duration_minutes=2,
    )
    question = QuestionDraft(
        question_id="question-guides",
        lesson_id=lesson.lesson_id,
        kind="single_choice",
        prompt="Какие направляющие используются в коллекции Чикаго Стрит?",
        options=(
            "Роликовые направляющие: плавный бесшумный ход ящиков.",
            "Роликовые направляющие на комодах: плавный ровный ход.",
            "Шариковые направляющие полного выдвижения.",
        ),
        correct_answer="Роликовые направляющие: плавный бесшумный ход ящиков.",
        explanation="Ответ подтверждён исходным материалом.",
        fact_id="fact-guides",
    )

    report = evaluate_publishability(
        course=CourseDraft(title="Коллекции", description="", lessons=(lesson,)),
        assessment=AssessmentDraft(questions=(question,)),
        blocks=(GroundedBlock(
            lesson_id=lesson.lesson_id,
            heading="Направляющие",
            text=lesson.content,
            fact_ids=lesson.fact_ids,
        ),),
        planned_fact_ids={"fact-guides"},
        provider_fallback_count=0,
    )

    assert report.publishable is False
    assert "unprofessional_learner_language" in report.reasons
    assert "ambiguous_question_options" in report.reasons


def test_v2_rejects_material_question_when_distractor_repeats_the_material_family() -> None:
    lesson = LessonDraft(
        lesson_id="lesson-material",
        module_title="Коллекции",
        title="Материалы коллекции",
        objective="Различать материалы",
        content="Корпус выполнен из ЛДСП Kronospan.",
        fact_ids=("fact-material",),
        supporting_fact_ids=(),
        duration_minutes=2,
    )
    question = QuestionDraft(
        question_id="question-material",
        lesson_id=lesson.lesson_id,
        kind="single_choice",
        prompt="Из какого материала выполнен корпус «Чикаго Нео»?",
        options=(
            "ЛДСП Kronospan (Австрия). Задние стенки и донья ящиков — ЛХДФ Kronospan.",
            "ЛДСП Kronospan (Австрия). Задние стенки — ЛХДФ Kronospan.",
            "Гладкий МДФ без ручек.",
        ),
        correct_answer=(
            "ЛДСП Kronospan (Австрия). Задние стенки и донья ящиков — ЛХДФ Kronospan."
        ),
        explanation="Корпус выполнен из ЛДСП Kronospan.",
        fact_id="fact-material",
    )

    report = evaluate_publishability(
        course=CourseDraft(title="Коллекции", description="", lessons=(lesson,)),
        assessment=AssessmentDraft(questions=(question,)),
        blocks=(
            GroundedBlock(
                lesson_id=lesson.lesson_id,
                heading="Материалы",
                text=lesson.content,
                fact_ids=lesson.fact_ids,
            ),
        ),
        planned_fact_ids={"fact-material"},
        provider_fallback_count=0,
    )

    assert report.publishable is False
    assert "ambiguous_question_options" in report.reasons


def test_v2_material_ambiguity_covers_plural_russian_and_english_prompts() -> None:
    for prompt in (
        "Из каких материалов выполнен фасад?",
        "Какие материалы используются в фасаде?",
        "Which materials are used in the facade?",
        "What is the facade made of?",
    ):
        question = QuestionDraft(
            question_id="question-material-plural",
            lesson_id="lesson-material",
            kind="single_choice",
            prompt=prompt,
            options=("MDF and metal", "MDF with glass", "solid wood"),
            correct_answer="MDF and metal",
            explanation="The source specifies MDF and metal.",
            fact_id="fact-material",
        )

        assert question_has_ambiguous_options(question) is True


def test_v2_keeps_material_question_with_distinct_material_families() -> None:
    question = QuestionDraft(
        question_id="question-material-distinct",
        lesson_id="lesson-material",
        kind="single_choice",
        prompt="Из какого материала выполнен фасад?",
        options=("МДФ", "стекло", "металл"),
        correct_answer="МДФ",
        explanation="Фасад выполнен из МДФ.",
        fact_id="fact-material",
    )

    assert question_has_ambiguous_options(question) is False


def test_v2_rejects_distractor_that_extends_short_correct_answer() -> None:
    question = QuestionDraft(
        question_id="question-scenario",
        lesson_id="lesson-scenario",
        kind="single_choice",
        prompt="Какое решение следует применить в ситуации клиента?",
        options=("Альта", "Нордик", "Терра", "Терра с разъяснением"),
        correct_answer="Терра",
        explanation="Источник указывает коллекцию Терра.",
        fact_id="fact-scenario",
    )

    assert question_has_ambiguous_options(question) is True


def test_v2_rejects_short_distractor_contained_in_longer_correct_answer() -> None:
    question = QuestionDraft(
        question_id="question-scenario-qualified",
        lesson_id="lesson-scenario",
        kind="single_choice",
        prompt="Какое решение следует применить в ситуации клиента?",
        options=("Альта", "Нордик", "Терра", "Терра с разъяснением"),
        correct_answer="Терра с разъяснением",
        explanation="Источник требует выбрать Терра с разъяснением.",
        fact_id="fact-scenario-qualified",
    )

    assert question_has_ambiguous_options(question) is True


def test_v2_allows_distinct_words_that_only_share_a_character_substring() -> None:
    question = QuestionDraft(
        question_id="question-distinct-substring",
        lesson_id="lesson-room",
        kind="single_choice",
        prompt="Для какой комнаты предназначена коллекция?",
        options=("Столовая", "Спальня", "Прихожая", "Гостиная"),
        correct_answer="Столовая",
        explanation="Источник относит коллекцию к столовой.",
        fact_id="fact-room",
    )

    assert question_has_ambiguous_options(question) is False
