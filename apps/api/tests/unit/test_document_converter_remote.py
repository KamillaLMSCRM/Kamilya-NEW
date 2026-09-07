from __future__ import annotations

import logging

import httpx
import pytest
from docx import Document

from app.modules.ai.ingestion import DocumentConverter
from app.modules.documents.archive_preflight import ArchivePreflightError


class _Response:
    def raise_for_status(self) -> None:
        return None

    def json(self) -> dict:
        return {
            "markdown": "# Policy",
            "pages": 2,
            "tables": 1,
            "engine": "markitdown",
            "engine_version": "0.1.6",
            "fallback_used": True,
            "warnings": ["Primary converter failed"],
            "profile": "office",
            "routing_reason": "office-docling-fallback",
        }


class _Client:
    def __init__(self, *args, **kwargs):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return None

    async def post(self, *args, **kwargs) -> _Response:
        return _Response()


async def test_remote_converter_preserves_engine_metadata(tmp_path, monkeypatch) -> None:
    source = tmp_path / "policy.docx"
    document = Document()
    document.add_paragraph("Approved policy")
    document.save(source)
    monkeypatch.setattr("httpx.AsyncClient", _Client)

    converted = await DocumentConverter("https://converter.example").convert(str(source))

    assert converted["markdown"] == "# Policy"
    assert converted["metadata"] == {
        "filename": "policy.docx",
        "size": source.stat().st_size,
        "pages": 2,
        "tables": 1,
        "engine": "markitdown",
        "engine_version": "0.1.6",
        "fallback_used": True,
        "warnings": ["Primary converter failed"],
        "profile": "office",
        "routing_reason": "office-docling-fallback",
    }


@pytest.mark.asyncio
async def test_invalid_ooxml_is_rejected_before_remote_converter(
    tmp_path, monkeypatch
) -> None:
    source = tmp_path / "policy.docx"
    source.write_bytes(b"PK fixture")

    async def unexpected_post(*args, **kwargs):
        raise AssertionError("remote converter must not receive invalid OOXML")

    monkeypatch.setattr(_Client, "post", unexpected_post)
    monkeypatch.setattr("httpx.AsyncClient", _Client)

    with pytest.raises(ArchivePreflightError):
        await DocumentConverter("https://converter.example").convert(str(source))


@pytest.mark.asyncio
async def test_remote_http_failure_logs_only_status_and_uses_local_fallback(
    tmp_path, monkeypatch, caplog
) -> None:
    source = tmp_path / "policy.docx"
    document = Document()
    document.add_paragraph("Approved policy")
    document.save(source)

    class FailedResponse(_Response):
        def raise_for_status(self) -> None:
            request = httpx.Request("POST", "https://converter.example/convert")
            response = httpx.Response(503, request=request)
            raise httpx.HTTPStatusError("unavailable", request=request, response=response)

    class FailedClient(_Client):
        async def post(self, *args, **kwargs) -> FailedResponse:
            return FailedResponse()

    monkeypatch.setattr("httpx.AsyncClient", FailedClient)
    caplog.set_level(logging.WARNING)

    converted = await DocumentConverter("https://converter.example").convert(str(source))

    assert converted["metadata"]["engine"] == "python-docx"
    assert "error_type=HTTPStatusError status_code=503" in caplog.text
    assert "Approved policy" not in caplog.text
