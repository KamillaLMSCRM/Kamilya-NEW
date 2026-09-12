from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.modules.ai.direct_source import (
    DirectSourceChunk,
    DirectSourceCorpus,
    DirectSourceDocument,
    _merged_worksheet_tables,
)
from app.modules.ai.document_passport import SectionRole, build_document_passport
from app.modules.ai.ingestion import DocumentChunker, DocumentConverter


@pytest.mark.asyncio
async def test_xlsx_conversion_preserves_visible_worksheet_boundaries(tmp_path: Path) -> None:
    workbook = Workbook()
    collections = workbook.active
    collections.title = "Collections"
    collections.append(["Collection", "Benefit"])
    collections.append(["Chicago", "Modular storage"])
    catalog = workbook.create_sheet("SKU catalog")
    catalog.append(["SKU", "Item", "Price"])
    catalog.append(["SKU-1", "Wardrobe", 199000])
    hidden = workbook.create_sheet("Hidden calculations")
    hidden.sheet_state = "hidden"
    hidden.append(["internal", "formula"])
    source = tmp_path / "source.xlsx"
    workbook.save(source)

    converter = DocumentConverter()
    converter.base_url = ""
    converted = await converter.convert(str(source))

    assert "# [Worksheet] Collections" in converted["markdown"]
    assert "# [Worksheet] SKU catalog" in converted["markdown"]
    assert "Hidden calculations" not in converted["markdown"]
    assert [item["name"] for item in converted["metadata"]["worksheets"]] == [
        "Collections",
        "SKU catalog",
    ]
    assert converted["metadata"]["fallback_used"] is False
    chunks = DocumentChunker(chunk_size=120, chunk_overlap=0).chunk_markdown(
        converted["markdown"],
        "doc-1",
        "source.xlsx",
    )
    catalog_chunks = [chunk for chunk in chunks if "SKU-1" in chunk["text"]]
    assert catalog_chunks
    assert "SKU catalog" in catalog_chunks[0]["metadata"]["headings"]
    assert "Collections" not in catalog_chunks[0]["metadata"]["headings"]

    direct_chunks = tuple(
        DirectSourceChunk(
            chunk_id=f"direct:doc-1:{index}",
            doc_id="doc-1",
            doc_name="source.xlsx",
            title="Product source",
            headings=tuple(json.loads(chunk["metadata"]["headings"])),
            text=chunk["text"],
            source_revision="document:" + "a" * 64,
            chunk_index=index,
        )
        for index, chunk in enumerate(chunks)
    )
    passport = build_document_passport(
        DirectSourceCorpus(
            tenant_id="tenant-1",
            documents=(
                DirectSourceDocument(
                    doc_id="doc-1",
                    title="Product source",
                    filename="source.xlsx",
                    category="general",
                    source_revision="document:" + "a" * 64,
                    chunks=direct_chunks,
                ),
            ),
            total_chars=len(converted["markdown"]),
            total_chunks=len(direct_chunks),
        )
    )
    roles = {section.name: section.role for section in passport.sections}
    assert roles == {
        "Collections": SectionRole.PRIMARY,
        "SKU catalog": SectionRole.SUPPORTING,
    }


@pytest.mark.asyncio
async def test_xlsx_chunking_bounds_long_rows_and_preserves_column_context(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Collections"
    sheet.append(["Field", "Phoenix", "Chicago"])
    phoenix_value = "Phoenix material " + "A" * 240
    chicago_value = "Chicago material " + "B" * 240
    sheet.append(["Materials", phoenix_value, chicago_value])
    source = tmp_path / "long-comparison.xlsx"
    workbook.save(source)

    converter = DocumentConverter()
    converter.base_url = ""
    converted = await converter.convert(str(source))
    chunks = DocumentChunker(chunk_size=160, chunk_overlap=32).chunk_markdown(
        converted["markdown"],
        "doc-1",
        source.name,
    )

    row_chunks = [chunk for chunk in chunks if "Materials" in chunk["text"]]
    assert len(row_chunks) > 1
    assert all(len(chunk["text"]) <= 160 for chunk in chunks)
    assert all("| Field |" in chunk["text"] for chunk in row_chunks)
    assert all("| --- | --- |" in chunk["text"] for chunk in row_chunks)
    from app.modules.ai.assessment import _markdown_tables

    values: dict[str, str] = {}
    for chunk in row_chunks:
        for headers, rows in _markdown_tables(chunk["text"]):
            for cells, _raw_row in rows:
                for header, value in zip(headers[1:], cells[1:], strict=True):
                    values[header] = values.get(header, "") + value
    combined = "".join(chunk["text"] for chunk in row_chunks)
    assert "Phoenix material" in combined
    assert "Chicago material" in combined
    assert values == {"Phoenix": phoenix_value, "Chicago": chicago_value}


@pytest.mark.asyncio
async def test_pathological_xlsx_labels_remain_bounded_and_reassemblable(
    tmp_path: Path,
) -> None:
    subject_header = "H" * 400
    column_header = "C" * 400
    subject = "S" * 400
    value = "V" * 240
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Collections"
    sheet.append([subject_header, column_header])
    sheet.append([subject, value])
    source = tmp_path / "pathological-labels.xlsx"
    workbook.save(source)

    converter = DocumentConverter()
    converter.base_url = ""
    converted = await converter.convert(str(source))
    chunks = DocumentChunker(chunk_size=160, chunk_overlap=0).chunk_markdown(
        converted["markdown"],
        "doc-1",
        source.name,
    )

    assert all(len(chunk["text"]) <= 160 for chunk in chunks)
    direct_chunks = tuple(
        DirectSourceChunk(
            chunk_id=f"direct:doc-1:{index}",
            doc_id="doc-1",
            doc_name=source.name,
            title="Product source",
            headings=tuple(json.loads(chunk["metadata"]["headings"])),
            text=chunk["text"],
            source_revision="document:" + "d" * 64,
            chunk_index=index,
            table_fragment=chunk["metadata"].get("worksheet_table_fragment"),
        )
        for index, chunk in enumerate(chunks)
    )
    merged = _merged_worksheet_tables(direct_chunks, {("doc-1", "collections")})

    assert len(merged) == 1
    assert merged[0][3] == [subject_header, column_header]
    assert merged[0][4][0][0] == [subject, value]


@pytest.mark.asyncio
async def test_xlsx_conversion_rejects_rows_beyond_streaming_limit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.modules.ai import ingestion

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Collection", "Benefit"])
    sheet.append(["Chicago", "Modular storage"])
    sheet.append(["Jagger", "Compact storage"])
    source = tmp_path / "oversized.xlsx"
    workbook.save(source)
    monkeypatch.setattr(ingestion, "MAX_XLSX_ROWS_PER_SHEET", 2)

    converter = DocumentConverter()
    converter.base_url = ""
    with pytest.raises(RuntimeError, match="safe conversion limits"):
        await converter.convert(str(source))
